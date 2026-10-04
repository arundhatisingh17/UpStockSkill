"""Creates paper_ledger.xlsx: a simulated-trade ledger with weekly compounded returns.
No broker connection. Nothing here places real orders."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = "/Users/induranasingh/Desktop/paper-trading/paper_ledger.xlsx"
MAX_TRADES = 500
MAX_WEEKS = 12

wb = Workbook()
bold = Font(bold=True)
head = PatternFill("solid", fgColor="DDEBF7")
inp = PatternFill("solid", fgColor="FFF2CC")

# ---- Settings
s = wb.active
s.title = "Settings"
rows = [
    ("Starting capital ($)", 10000, "Editable. Assumed value."),
    ("Stock commission per share ($)", 0.0, "Schwab $0 online stock trades"),
    ("Options commission per contract ($)", 0.65, "Schwab standard; verify current"),
    ("First week start (Monday)", "2026-10-05", "Edit to the Monday your test begins"),
    ("Fill rule", "Buy at ask, sell at bid (+ slippage)", "Never fill at mid or last"),
]
s["A1"], s["B1"], s["C1"] = "Setting", "Value", "Note"
for c in "ABC":
    s[f"{c}1"].font = bold
    s[f"{c}1"].fill = head
for i, r in enumerate(rows, start=2):
    for j, v in enumerate(r):
        s.cell(i, j + 1, v)
    s.cell(i, 2).fill = inp
s.column_dimensions["A"].width = 38
s.column_dimensions["B"].width = 38
s.column_dimensions["C"].width = 48

# ---- Trades
t = wb.create_sheet("Trades")
cols = [
    "ID", "Signal source (skill)", "Ticker", "Instrument", "Side (Long/Short)", "Qty",
    "Multiplier", "Open date", "Open fill", "Close date", "Close fill", "Fees ($)",
    "Gross P/L ($)", "Net P/L ($)", "Week start", "Status", "Signal details / notes",
]
for j, c in enumerate(cols, start=1):
    cell = t.cell(1, j, c)
    cell.font = bold
    cell.fill = head
    cell.alignment = Alignment(wrap_text=True)
    t.column_dimensions[get_column_letter(j)].width = 16
t.column_dimensions["Q"].width = 50
t.freeze_panes = "A2"
for r in range(2, MAX_TRADES + 2):
    # Gross P/L: only when closed
    t[f"M{r}"] = (
        f'=IF(OR(I{r}="",K{r}=""),"",IF(E{r}="Short",-1,1)*(K{r}-I{r})*F{r}*IF(G{r}="",1,G{r}))'
    )
    t[f"N{r}"] = f'=IF(M{r}="","",M{r}-L{r})'
    t[f"O{r}"] = f'=IF(J{r}="","",DATEVALUE(J{r})-WEEKDAY(DATEVALUE(J{r}),3))'
    t[f"P{r}"] = f'=IF(I{r}="","",IF(K{r}="","Open","Closed"))'
    t[f"O{r}"].number_format = "yyyy-mm-dd"

# ---- Weekly
w = wb.create_sheet("Weekly")
wc = ["Week #", "Week start", "Start equity ($)", "Net P/L ($)", "End equity ($)",
      "Weekly return", "Cumulative (compounded)", "Trades closed"]
for j, c in enumerate(wc, start=1):
    cell = w.cell(1, j, c)
    cell.font = bold
    cell.fill = head
    w.column_dimensions[get_column_letter(j)].width = 22
for k in range(1, MAX_WEEKS + 1):
    r = k + 1
    w[f"A{r}"] = k
    w[f"B{r}"] = f'=DATEVALUE(Settings!$B$5)+7*(A{r}-1)'
    w[f"B{r}"].number_format = "yyyy-mm-dd"
    w[f"C{r}"] = "=Settings!$B$2" if k == 1 else f"=E{r-1}"
    w[f"D{r}"] = f"=SUMIFS(Trades!$N$2:$N${MAX_TRADES+1},Trades!$O$2:$O${MAX_TRADES+1},B{r})"
    w[f"E{r}"] = f"=C{r}+D{r}"
    w[f"F{r}"] = f"=IF(C{r}=0,0,D{r}/C{r})"
    w[f"G{r}"] = f"=E{r}/Settings!$B$2-1"
    w[f"H{r}"] = f"=COUNTIFS(Trades!$O$2:$O${MAX_TRADES+1},B{r})"
    for c in "CDE":
        w[f"{c}{r}"].number_format = "#,##0.00"
    for c in "FG":
        w[f"{c}{r}"].number_format = "0.00%"

# ---- Summary
m = wb.create_sheet("Summary")
N = f"Trades!$N$2:$N${MAX_TRADES+1}"
items = [
    ("Closed trades", f"=COUNT({N})", "0"),
    ("Net P/L ($)", f"=SUM({N})", "#,##0.00"),
    ("Total return (compounded)", f"=Weekly!E{MAX_WEEKS+1}/Settings!B2-1", "0.00%"),
    ("Average weekly return", "=AVERAGEIF(Weekly!H2:H13,\">0\",Weekly!F2:F13)", "0.00%"),
    ("Geometric weekly return", f"=(Weekly!E{MAX_WEEKS+1}/Settings!B2)^(1/MAX(1,COUNTIF(Weekly!H2:H13,\">0\")))-1", "0.00%"),
    ("Win rate", f'=IF(COUNT({N})=0,0,COUNTIF({N},">0")/COUNT({N}))', "0.0%"),
    ("Average win ($)", f'=IFERROR(AVERAGEIF({N},">0"),0)', "#,##0.00"),
    ("Average loss ($)", f'=IFERROR(AVERAGEIF({N},"<0"),0)', "#,##0.00"),
    ("Expectancy per trade ($)", f'=IF(COUNT({N})=0,0,SUM({N})/COUNT({N}))', "#,##0.00"),
    ("Largest loss ($)", f"=MIN({N})", "#,##0.00"),
    ("Open positions", f'=COUNTIF(Trades!$P$2:$P${MAX_TRADES+1},"Open")', "0"),
]
m["A1"], m["B1"] = "Metric", "Value"
for c in "AB":
    m[f"{c}1"].font = bold
    m[f"{c}1"].fill = head
for i, (a, f, nf) in enumerate(items, start=2):
    m[f"A{i}"], m[f"B{i}"] = a, f
    m[f"B{i}"].number_format = nf
n = len(items) + 3
m[f"A{n}"] = "Reality check"
m[f"A{n}"].font = bold
m[f"A{n+1}"] = '=IF(B2<100,"Only "&B2&" closed trades. Under 100 is too few to trust the results.","Sample size OK. Still check costs, drawdown, and regime.")'
m[f"A{n+2}"] = "Fills assume ask/bid + slippage on delayed quotes; real fills can be worse. Paper results are optimistic."
m.column_dimensions["A"].width = 34
m.column_dimensions["B"].width = 18

wb.move_sheet("Summary", offset=-3)
wb.save(OUT)
print("saved", OUT)
